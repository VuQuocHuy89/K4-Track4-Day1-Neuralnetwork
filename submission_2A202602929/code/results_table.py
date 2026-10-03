"""Lưu JSON thí nghiệm và tạo workbook kết quả từ template.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu được giữ nguyên)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import torch


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    cfg = result.get("cfg", {})
    exp_id = cfg.get("exp_id")
    if not exp_id:
        raise ValueError("Mỗi kết quả cần cfg['exp_id'] để lưu JSON.")
    output_dir = Path(results_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "cfg": cfg,
        "history": result.get("history", {}),
        "summary": result.get("summary", {}),
    }
    path = output_dir / f"{exp_id}.json"
    with path.open("w", encoding="utf-8") as stream:
        json.dump(_to_jsonable(payload), stream, ensure_ascii=False, indent=2, allow_nan=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    root = Path(results_dir)
    if not root.exists():
        return []
    results = []
    for path in sorted(root.glob("*.json")):
        with path.open(encoding="utf-8") as stream:
            result = json.load(stream)
        if "cfg" in result and "history" in result and "summary" in result:
            results.append(result)
    return sorted(results, key=lambda result: result["cfg"].get("exp_id", ""))


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")
    hidden = cfg.get("hidden", ())
    hidden_text = "-".join(str(width) for width in hidden) if hidden else ""
    optimizer_labels = {
        "sgd": "SGD",
        "sgd_momentum": "SGD+momentum",
        "adam": "Adam",
        "adamw": "AdamW",
    }
    loss_labels = {"ce": "CE", "mse": "MSE"}
    notes_text = notes
    if (
        cfg.get("optimizer") in {"sgd", "sgd_momentum", "adam", "adamw"}
        or cfg.get("scheduler")
        or cfg.get("clip_norm") is not None
    ):
        extra = []
        if cfg.get("optimizer") == "sgd_momentum":
            extra.append(f"momentum={cfg.get('momentum', 0.9)}")
        if cfg.get("optimizer") in {"adam", "adamw"}:
            extra.append(f"betas={cfg.get('betas', (0.9, 0.999))}")
            extra.append(f"eps={cfg.get('eps', 1e-8)}")
        if cfg.get("scheduler"):
            extra.append(f"scheduler={cfg['scheduler']}")
        if cfg.get("clip_norm") is not None:
            fraction = summary.get("clip_fraction")
            extra.append(
                f"clip_norm={cfg['clip_norm']}"
                if fraction is None
                else f"clip_norm={cfg['clip_norm']}; clipped_steps={fraction:.1%}"
            )
        config_note = "; ".join(extra)
        notes_text = f"{notes_text}; {config_note}" if notes_text else config_note

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": loss_labels.get(cfg.get("loss"), cfg.get("loss", "")),
        "optimizer": optimizer_labels.get(cfg.get("optimizer"), cfg.get("optimizer", "")),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch"),
        "epochs": cfg.get("epochs"),
        "hidden": hidden_text,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": "none" if cfg.get("clip_norm") is None else cfg.get("clip_norm"),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed"),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": bool(summary.get("diverged", False)),
        "eval_acc": None if eval_scores is None else eval_scores.get("accuracy"),
        "eval_macro_f1": None if eval_scores is None else eval_scores.get("macro_f1"),
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes_text,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Ghi kết quả vào sheet Experiments, giữ nguyên các cột công thức của template."""
    try:
        from copy import copy
        from openpyxl import load_workbook
        from openpyxl.formula.translate import Translator
    except ImportError as exc:
        raise ImportError("write_xlsx cần thư viện openpyxl.") from exc

    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError(f"Không tìm thấy template Excel: {template}")
    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template)
    if "Experiments" not in workbook.sheetnames:
        raise KeyError("Template thiếu sheet 'Experiments'.")
    worksheet = workbook["Experiments"]
    headers = [cell.value for cell in worksheet[1]]
    if not headers or headers[0] != "exp_id":
        raise ValueError("Không đọc được header của sheet Experiments.")

    formula_columns = {
        "step0_gap_vs_lnC",
        "gap_val_minus_train",
        "delta_val_f1_vs_base",
        "beyond_noise",
    }
    for row_index, row in enumerate(rows, start=2):
        if row_index > worksheet.max_row:
            source_row = 2
            worksheet.row_dimensions[row_index].height = worksheet.row_dimensions[source_row].height
            for column_index in range(1, worksheet.max_column + 1):
                source = worksheet.cell(source_row, column_index)
                target = worksheet.cell(row_index, column_index)
                if source.has_style:
                    target._style = copy(source._style)
                if source.number_format:
                    target.number_format = source.number_format
                if source.data_type == "f":
                    target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)

        for column_index, header in enumerate(headers, start=1):
            if header in formula_columns or header is None:
                continue
            worksheet.cell(row_index, column_index).value = _to_excel_value(row.get(header))

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass
    workbook.save(output)


def _to_jsonable(value):
    """Đổi kiểu NumPy/PyTorch và NaN/inf sang giá trị JSON hợp lệ."""
    if isinstance(value, dict):
        return {str(key): _to_jsonable(item) for key, item in value.items() if key != "best_state"}
    if isinstance(value, (list, tuple)):
        return [_to_jsonable(item) for item in value]
    if isinstance(value, torch.Tensor):
        return _to_jsonable(value.detach().cpu().tolist())
    if isinstance(value, np.ndarray):
        return _to_jsonable(value.tolist())
    if isinstance(value, np.generic):
        return _to_jsonable(value.item())
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _to_excel_value(value):
    if value is None:
        return None
    if isinstance(value, (tuple, list)):
        return "-".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, torch.Tensor):
        return value.detach().cpu().item()
    return value
